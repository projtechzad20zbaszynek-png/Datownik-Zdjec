import os,re,csv,threading,subprocess,sys
from pathlib import Path
from datetime import datetime,date
import tkinter as tk
from tkinter import ttk,filedialog,messagebox
from PIL import Image,ImageTk,ImageDraw,ImageFont,ImageOps
from openpyxl import load_workbook
EXT={".jpg",".jpeg",".png",".tif",".tiff"}; DATE_TAGS=(306,36867,36868); GPS_TAG=34853
def norm(v):
 s=str(v or "").strip().upper().replace("/", "-").replace("\\","-")
 return re.sub("-+","-",re.sub(r"\s+"," ",s))
def pdate(v):
 if isinstance(v,datetime):return v.date()
 if isinstance(v,date):return v
 s=str(v or "").strip().split(" ")[0]
 for f in ("%d.%m.%Y","%Y-%m-%d","%d-%m-%Y","%d/%m/%Y"):
  try:return datetime.strptime(s,f).date()
  except:pass
def load_dates(files):
 out={};conf=[]
 for fn in files:
  wb=load_workbook(fn,data_only=True,read_only=True)
  for ws in wb.worksheets:
   for hr in range(1,min(ws.max_row,30)+1):
    vals=[str(c.value or "").strip() for c in ws[hr]]
    dc=next((i+1 for i,v in enumerate(vals) if v.casefold()=="data wykonania kontroli"),None)
    ids=[i+1 for i,v in enumerate(vals) if v.casefold()=="id_bud"]
    if not dc or not ids:continue
    ic=1 if 1 in ids else ids[0]
    for r in range(hr+1,ws.max_row+1):
     k=norm(ws.cell(r,ic).value);d=pdate(ws.cell(r,dc).value)
     if k and d:
      if k in out and out[k]!=d:conf.append(k)
      else:out[k]=d
    break
  wb.close()
 return out,sorted(set(conf))
def xtime(ex):
 for t in (36867,36868,306):
  v=ex.get(t)
  if v:
   try:return datetime.strptime(str(v),"%Y:%m:%d %H:%M:%S").time()
   except:pass
def gps(ex):
 try:
  g=ex.get_ifd(GPS_TAG)
  def dms(v):return float(v[0])+float(v[1])/60+float(v[2])/3600
  a=dms(g[2]);b=dms(g[4])
  if g[1] in ("S",b"S"):a=-a
  if g[3] in ("W",b"W"):b=-b
  return a,b
 except:return None
STYLES={
 "LCD 7-segment":("segment",(255,145,25),(55,18,0)),
 "OSD kamera":("osd",(255,255,255),(0,0,0)),
 "Pixel retro":("pixel",(255,175,25),(45,15,0)),
 "Pomarańczowy klasyczny":("classic",(255,125,20),(65,20,0)),
 "Żółty retro":("retro",(255,220,35),(45,35,0))
}
SEG={"0":"abcedf","1":"bc","2":"abdeg","3":"abcdg","4":"bcfg","5":"acdfg","6":"acdefg","7":"abc","8":"abcdefg","9":"abcdfg"}
def style_font(kind,size):
 names={"osd":"lucon.ttf","pixel":"courbd.ttf","classic":"arialbd.ttf","retro":"consolab.ttf"}
 p=os.path.join(os.environ.get("WINDIR","C:/Windows"),"Fonts",names.get(kind,"consolab.ttf"))
 try:return ImageFont.truetype(p,size)
 except:return ImageFont.load_default()
def draw_segment_text(im,pos,text,size,fill,stroke,anchor="nw"):
 scale=max(1,size//10);cw=int(size*.62);ch=size;gap=max(2,size//7);total=len(text)*(cw+gap)-gap
 x,y=pos
 if "e" in anchor:x-=total
 elif anchor=="center":x-=total/2
 if "s" in anchor:y-=ch
 elif anchor=="center":y-=ch/2
 d=ImageDraw.Draw(im)
 segs={"a":(.16,.05,.84,.05),"b":(.88,.12,.88,.46),"c":(.88,.54,.88,.88),"d":(.16,.95,.84,.95),"e":(.08,.54,.08,.88),"f":(.08,.12,.08,.46),"g":(.16,.50,.84,.50)}
 for c in text:
  if c.isdigit():
   for k in SEG.get(c,""):
    q=segs[k];d.line((x+q[0]*cw,y+q[1]*ch,x+q[2]*cw,y+q[3]*ch),fill=stroke,width=scale*3)
    d.line((x+q[0]*cw,y+q[1]*ch,x+q[2]*cw,y+q[3]*ch),fill=fill,width=scale)
  elif c in ".-/:":
   if c==".":d.ellipse((x+cw*.42,y+ch*.84,x+cw*.58,y+ch),fill=fill)
   elif c==":":
    d.ellipse((x+cw*.42,y+ch*.28,x+cw*.58,y+ch*.44),fill=fill);d.ellipse((x+cw*.42,y+ch*.65,x+cw*.58,y+ch*.81),fill=fill)
   else:d.line((x+cw*.18,y+ch*(.72 if c=="/" else .5),x+cw*.82,y+ch*(.28 if c=="/" else .5)),fill=fill,width=scale*2)
  x+=cw+gap
 return total,ch
def draw_stamp(im,pos,text,style,size,anchor):
 kind,fill,stroke=STYLES[style]
 if kind=="segment":return draw_segment_text(im,pos,text,size,fill,stroke,anchor)
 d=ImageDraw.Draw(im);fo=style_font(kind,size);sw=max(1,size//18)
 if kind=="pixel":sw=max(1,size//24)
 bb=d.textbbox((0,0),text,font=fo,stroke_width=sw);tw=bb[2]-bb[0];th=bb[3]-bb[1];x,y=pos
 off={"nw":(0,0),"ne":(tw,0),"sw":(0,th),"se":(tw,th),"center":(tw/2,th/2)}.get(anchor,(0,0))
 d.text((x-off[0],y-off[1]),text,font=fo,fill=fill,stroke_width=sw,stroke_fill=stroke)
 return tw,th
def textval(d,layout,sep,ex,tm,gp):
 s=(f"{d.day:02d}{sep}{d.month:02d}{sep}{d.year:04d}" if layout=="DD MM RRRR" else f"{d.year:04d}{sep}{d.month:02d}{sep}{d.day:02d}")
 a=[s]
 if tm:
  t=xtime(ex)
  if t:a.append(t.strftime("%H:%M:%S"))
 if gp:
  g=gps(ex)
  if g:a.append(f"{g[0]:.5f}°, {g[1]:.5f}°")
 return "  ".join(a)
class App(tk.Tk):
 def __init__(self):
  super().__init__();self.title("Datownik by Projekt 277 PJ");self.geometry("1240x820");self.minsize(1080,720);self.configure(bg="#0b3d73")
  try:
   ico=Path(getattr(sys,"_MEIPASS",Path(__file__).parent))/"P277.ico";self.iconbitmap(str(ico))
  except:pass
  st=ttk.Style(self);st.theme_use("clam");st.configure(".",background="#0b3d73",foreground="white",fieldbackground="#eaf4ff");st.configure("TFrame",background="#0b3d73");st.configure("TLabel",background="#0b3d73",foreground="white");st.configure("TCheckbutton",background="#0b3d73",foreground="white");st.configure("TButton",background="#1677d2",foreground="white",padding=6);st.map("TButton",background=[("active","#2493f2")]);st.configure("TCombobox",fieldbackground="white",foreground="#082b50")
  self.src=tk.StringVar();self.exc=[];self.preview=None;self.photo=None;self.box=None;self.preview_img=None;self.preview_exif={};self.date_map={};self.date_conf=[]
  self.fontv=tk.StringVar(value="LCD 7-segment");self.layout=tk.StringVar(value="DD MM RRRR");self.sep=tk.StringVar(value=".")
  self.sz=tk.IntVar(value=44);self.tm=tk.BooleanVar();self.gp=tk.BooleanVar();self.exif=tk.BooleanVar(value=True)
  self.manual=tk.BooleanVar();self.mdate=tk.StringVar(value=datetime.now().strftime("%d.%m.%Y"));self.sample_date=date(2026,6,10)
  self.x=.97;self.y=.95;self.anchor="se";self.ui()
 def ui(self):
  l=ttk.Frame(self,padding=12);l.pack(side="left",fill="y");r=ttk.Frame(self,padding=12);r.pack(fill="both",expand=True)
  ttk.Label(l,text="P277   Datownik by Projekt 277 PJ",font=("Segoe UI",18,"bold")).pack(anchor="w",pady=(0,2));ttk.Label(l,text="Projekt 277 - Patryk Jabłoński",font=("Segoe UI",9)).pack(anchor="w",pady=(0,10))
  ttk.Button(l,text="1. Wybierz folder zdjęć",command=self.pickfolder).pack(fill="x");self.sl=ttk.Label(l,text="Nie wybrano",wraplength=350);self.sl.pack(anchor="w",pady=(3,8))
  ttk.Button(l,text="2. Dodaj pliki Excel (opcjonalnie)",command=self.pickexcel).pack(fill="x");self.el=ttk.Label(l,text="Excel: 0");self.el.pack(anchor="w",pady=(3,4))
  ttk.Checkbutton(l,text="Użyj ręcznie wpisanej daty dla wszystkich zdjęć",variable=self.manual,command=self.refresh).pack(anchor="w")
  q=ttk.Frame(l);q.pack(fill="x",pady=(2,8));ttk.Label(q,text="Data:").pack(side="left");ttk.Entry(q,textvariable=self.mdate,width=14).pack(side="left",padx=5);ttk.Label(q,text="np. 10.06.2026").pack(side="left")
  for label,var,vals in [("Styl aparatu",self.fontv,list(STYLES.keys())),("Układ daty",self.layout,["DD MM RRRR","RRRR MM DD"]),("Separator",self.sep,[".","-","/"])]:
   ttk.Label(l,text=label).pack(anchor="w");c=ttk.Combobox(l,textvariable=var,values=vals,state="readonly");c.pack(fill="x",pady=(0,5));c.bind("<<ComboboxSelected>>",lambda e:self.refresh())
  ttk.Label(l,text="Rozmiar").pack(anchor="w");ttk.Scale(l,from_=20,to=90,variable=self.sz,command=lambda e:self.refresh()).pack(fill="x")
  ttk.Checkbutton(l,text="Dodaj godzinę z EXIF",variable=self.tm,command=self.refresh).pack(anchor="w")
  ttk.Checkbutton(l,text="Dodaj GPS z EXIF",variable=self.gp,command=self.refresh).pack(anchor="w")
  ttk.Checkbutton(l,text="Zmień datę wykonania EXIF (zachowaj godzinę)",variable=self.exif).pack(anchor="w")
  ttk.Label(l,text="Położenie").pack(anchor="w",pady=(8,2));p=ttk.Frame(l);p.pack(fill="x")
  for t,x,y,a in [("↖",.03,.05,"nw"),("↗",.97,.05,"ne"),("↙",.03,.95,"sw"),("↘",.97,.95,"se")]:ttk.Button(p,text=t,command=lambda x=x,y=y,a=a:self.setpos(x,y,a)).pack(side="left",expand=True,fill="x")
  ttk.Label(l,text="Podgląd pokazuje przykładową datę 10.06.2026. Ustaw jej wygląd i przeciągnij ją myszką. Podczas datowania program wstawi właściwą datę z Excel lub wpisaną ręcznie.",wraplength=350).pack(anchor="w",pady=7)
  ttk.Button(l,text="URUCHOM DATOWANIE",command=self.start).pack(fill="x",pady=(10,4));self.pb=ttk.Progressbar(l);self.pb.pack(fill="x");self.st=ttk.Label(l,text="Gotowy");self.st.pack(anchor="w")
  ttk.Label(r,text="Podgląd ustawienia datownika — przykład 10.06.2026",font=("Segoe UI",14,"bold")).pack(anchor="w");self.cv=tk.Canvas(r,bg="#06294d",highlightthickness=0);self.cv.pack(fill="both",expand=True,pady=(8,0));self.cv.bind("<Configure>",lambda e:self.refresh());self.cv.bind("<Button-1>",self.drag);self.cv.bind("<B1-Motion>",self.drag)
 def pickfolder(self):
  p=filedialog.askdirectory()
  if p:self.src.set(p);self.sl.config(text=p);self.first()
 def pickexcel(self):
  f=filedialog.askopenfilenames(filetypes=[("Excel","*.xlsx *.xlsm")])
  if f:
   self.exc=list(f)
   self.el.config(text=f"Excel: {len(f)}")
   self.st.config(text="Wczytuję Excel...")
   self.update_idletasks()
   try:
    self.date_map,self.date_conf=load_dates(self.exc)
    self.st.config(text=f"Wczytano {len(self.date_map)} dat")
   except Exception as e:
    self.date_map={};self.date_conf=[]
    messagebox.showerror("Excel",str(e))
   self.refresh()
 def first(self):
  self.preview=None;self.preview_img=None;self.preview_exif={}
  for dp,_,fs in os.walk(self.src.get()):
   for f in sorted(fs):
    if Path(f).suffix.lower() in EXT:
     self.preview=Path(dp)/f
     try:
      with Image.open(self.preview) as z:
       self.preview_exif=z.getexif()
       self.preview_img=ImageOps.exif_transpose(z).convert("RGB")
     except Exception as e:
      self.st.config(text="Podgląd: "+str(e))
     self.refresh()
     return
 def match(self,p,mp):
  cur=p.parent;root=Path(self.src.get())
  while True:
   k=norm(cur.name)
   if k in mp:return mp[k],k
   if cur==root or root not in cur.parents:break
   cur=cur.parent
  return None,""
 def chosen_manual(self):
  return pdate(self.mdate.get())
 def previewdate(self):
  # Podgląd zawsze używa stałej daty przykładowej.
  # Właściwe daty są pobierane dopiero podczas przetwarzania.
  return self.sample_date
 def refresh(self):
  if not hasattr(self,"cv"):return
  self.cv.delete("all")
  if not self.preview or self.preview_img is None:
   self.cv.create_text(max(200,self.cv.winfo_width()//2),max(120,self.cv.winfo_height()//2),text="Wybierz folder zdjęć — tutaj pojawi się podgląd datownika 10.06.2026",fill="white",font=("Segoe UI",16,"bold"),width=520)
   return
  try:
   im=self.preview_img.copy();ex=self.preview_exif
   # Najpierw nanosimy datę NA KOPIĘ PEŁNEGO ZDJĘCIA, dokładnie jak przy zapisie finalnym.
   txt=textval(self.sample_date,self.layout.get(),self.sep.get(),ex,self.tm.get(),self.gp.get())
   full_fs=max(18,int(min(im.width,im.height)*(self.sz.get()/1000.0)))
   draw_stamp(im,(self.x*im.width,self.y*im.height),txt,self.fontv.get(),full_fs,self.anchor)
   # Dopiero gotowe zdjęcie z datą skalujemy do okna podglądu.
   cw=max(100,self.cv.winfo_width());ch=max(100,self.cv.winfo_height());sc=min(cw/im.width,ch/im.height)
   w=max(1,int(im.width*sc));h=max(1,int(im.height*sc));ox=(cw-w)//2;oy=(ch-h)//2
   shown=im.resize((w,h),Image.Resampling.LANCZOS)
   self.photo=ImageTk.PhotoImage(shown);self.cv.create_image(ox,oy,image=self.photo,anchor="nw")
   self.box=(ox,oy,w,h);self.st.config(text="Podgląd: przykładowa data 10.06.2026")
  except Exception as e:
   self.st.config(text="Błąd podglądu: "+str(e))
   self.cv.create_text(20,20,text="Błąd podglądu: "+str(e),fill="white",anchor="nw",width=max(200,self.cv.winfo_width()-40))
 def drag(self,e):
  if self.box:
   ox,oy,w,h=self.box
   self.x=min(.99,max(.01,(e.x-ox)/w));self.y=min(.99,max(.01,(e.y-oy)/h));self.anchor="center";self.refresh()
 def setpos(self,x,y,a):
  self.x=x;self.y=y;self.anchor=a;self.refresh()
 def start(self):
  if not self.src.get():return messagebox.showwarning("Brak folderu","Wybierz folder zdjęć.")
  if self.manual.get():
   if not self.chosen_manual():return messagebox.showwarning("Błędna data","Wpisz datę, np. 10.06.2026.")
  elif not self.exc:return messagebox.showwarning("Brak daty","Dodaj Excel albo zaznacz ręczną datę.")
  threading.Thread(target=self.run,daemon=True).start()
 def run(self):
  try:
   mp={};conf=[]
   if not self.manual.get():
     mp,conf=(self.date_map,self.date_conf) if self.date_map else load_dates(self.exc)
   if conf:raise ValueError("Sprzeczne daty w Excel dla: "+", ".join(conf[:20]))
   md=self.chosen_manual() if self.manual.get() else None;root=Path(self.src.get());out=root.parent/(root.name+"_DATOWANE");files=[p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in EXT];self.pb["maximum"]=max(1,len(files))
   rows=[];missing=set();errors=[];ok=0
   for i,p in enumerate(files,1):
    if md:d,k=md,norm(p.parent.name)
    else:
     d,k=self.match(p,mp)
     if not d:
      # report nearest BUD folder name, else direct parent
      bud=next((a.name for a in [p.parent,*p.parents] if norm(a.name).startswith("BUD ")),p.parent.name);missing.add(bud);rows.append([str(p.relative_to(root)),bud,"BRAK DATY"]);self.after(0,lambda v=i:self.pb.configure(value=v));continue
    try:
     with Image.open(p) as src:ex=src.getexif();ot=xtime(ex);icc=src.info.get("icc_profile");im=ImageOps.exif_transpose(src).convert("RGB")
     txt=textval(d,self.layout.get(),self.sep.get(),ex,self.tm.get(),self.gp.get());font_size=max(16,int(min(im.width,im.height)*(self.sz.get()/1000.0)));draw_stamp(im,(self.x*im.width,self.y*im.height),txt,self.fontv.get(),font_size,self.anchor)
     if self.exif.get():
      val=f"{d:%Y:%m:%d} {ot:%H:%M:%S}" if ot else f"{d:%Y:%m:%d} 00:00:00"
      for t in DATE_TAGS:ex[t]=val
      ex[274]=1
     dest=out/p.relative_to(root);dest.parent.mkdir(parents=True,exist_ok=True);kw={}
     if dest.suffix.lower() in (".jpg",".jpeg"):kw={"quality":95,"subsampling":0}
     if ex:kw["exif"]=ex.tobytes()
     if icc:kw["icc_profile"]=icc
     im.save(dest,**kw);ok+=1;rows.append([str(p.relative_to(root)),k,d.strftime("%d.%m.%Y"),"OK"])
    except Exception as e:errors.append(f"{p}: {e}");rows.append([str(p.relative_to(root)),k,"BŁĄD"])
    self.after(0,lambda v=i,n=len(files):(self.pb.configure(value=v),self.st.config(text=f"{v}/{n}")))
   out.mkdir(parents=True,exist_ok=True);report=out/"Raport_Datownika.txt"
   with open(report,"w",encoding="utf-8-sig") as f:
    f.write("RAPORT - DATOWNIK ZDJĘĆ\n"+"="*60+"\n");f.write(f"Data raportu: {datetime.now():%d.%m.%Y %H:%M:%S}\nZdjęcia znalezione: {len(files)}\nZdjęcia przetworzone: {ok}\nFoldery bez przypisanej daty: {len(missing)}\nBłędy: {len(errors)}\n\n")
    f.write("FOLDERY, KTÓRYM NIE PRZYPISANO DATY\n"+"-"*60+"\n")
    if missing:
     for x in sorted(missing):f.write(x+"\n")
    else:f.write("Brak - wszystkim folderom przypisano datę.\n")
    if errors:
     f.write("\nBŁĘDY\n"+"-"*60+"\n")
     for x in errors:f.write(x+"\n")
   self.st.config(text=f"Gotowe {ok}/{len(files)}")
   try:os.startfile(report)
   except:pass
   messagebox.showinfo("Gotowe",f"Przetworzono: {ok}/{len(files)}\nRaport: {report}")
  except Exception as e:self.after(0,lambda err=str(e):messagebox.showerror("Błąd",err))
if __name__=="__main__":App().mainloop()
